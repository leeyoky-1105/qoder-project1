"""基金产品文件导入服务."""
import logging
from datetime import datetime

from openpyxl import load_workbook

from models import db
from models.fund import FundProduct

logger = logging.getLogger(__name__)


class FundLoaderService:
    """基金产品导入服务.

    从 Excel 配置文件加载基金产品信息到数据库.
    """

    @staticmethod
    def import_funds(config_path: str) -> dict:
        """从 Excel 文件导入基金产品.

        Args:
            config_path: Excel 配置文件路径.

        Returns:
            dict: 导入结果，包含成功/失败数量.
        """
        try:
            wb = load_workbook(config_path, read_only=True)
            ws = wb.active
        except Exception as e:
            logger.error("读取产品配置文件失败: %s", e)
            return {
                "code": 400,
                "message": f"配置文件读取失败: {e}",
                "data": None,
            }

        success_count = 0
        fail_count = 0
        errors = []

        rows = list(ws.iter_rows(min_row=2, values_only=True))
        wb.close()

        for row in rows:
            try:
                if not row or not row[0]:
                    continue

                fund_name = str(row[0]).strip()
                fund_code = str(row[1]).strip()
                nav = float(row[2])
                nav_date_raw = str(row[3]).strip()

                # 处理净值日期格式 (20260915 -> 2026-09-15)
                if len(nav_date_raw) == 8 and nav_date_raw.isdigit():
                    nav_date = (
                        f"{nav_date_raw[:4]}-"
                        f"{nav_date_raw[4:6]}-"
                        f"{nav_date_raw[6:8]}"
                    )
                else:
                    nav_date = nav_date_raw

                existing = FundProduct.query.filter_by(
                    fund_code=fund_code
                ).first()

                if existing:
                    existing.fund_name = fund_name
                    existing.nav = nav
                    existing.nav_date = nav_date
                else:
                    fund = FundProduct(
                        fund_code=fund_code,
                        fund_name=fund_name,
                        nav=nav,
                        nav_date=nav_date,
                        status="ACTIVE",
                    )
                    db.session.add(fund)

                success_count += 1
            except Exception as e:
                logger.warning(
                    "导入基金行失败: %s, 错误: %s",
                    row, e,
                )
                fail_count += 1
                errors.append(str(e))

        db.session.commit()
        logger.info(
            "产品导入完成: 成功 %d, 失败 %d",
            success_count,
            fail_count,
        )

        return {
            "code": 0,
            "message": "success",
            "data": {
                "success_count": success_count,
                "fail_count": fail_count,
                "errors": errors,
            },
        }
